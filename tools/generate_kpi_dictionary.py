from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "KPI_Dictionary.xlsx"

SOURCE_FIELDS = [
    ("Credit union name", "No", "name", "Passed through as the institution identifier.", "Identifies the reporting credit union and supports search, labels, and record matching."),
    ("City", "No", "city", "Passed through from the source; all 4,336 workbook rows are blank in this field.", "Could support local-market analysis, but no usable city values are present."),
    ("Headquarters state", "No", "state", "Passed through as the state code; used for state-market filtering and aggregation.", "Defines headquarters-state geography, not branch service area."),
    ("Region", "No", "region", "Passed through from the source.", "Provides a broad geographic grouping where supplied."),
    ("Asset bucket", "Yes", "bucket; assetsM", "Uses source bucket when a valid integer 0-7 is supplied; otherwise: <$50M=0, <$100M=1, <$250M=2, <$500M=3, <$1B=4, <$5B=5, <$10B=6, otherwise 7.", "Normalizes institutions into size cohorts. Current workbook provides this value directly."),
    ("Total assets ($M)", "No", "assetsM", "Source balance in $ millions; import normalization may divide dollar or thousand-denominated amounts by 1,000,000 or 1,000 based on median positive assets.", "Primary balance-sheet size measure used for ranking, market concentration, size tier, and peer shortlist."),
    ("Members", "No", "members", "Passed through as a reported member count.", "Measures membership scale; not population penetration without an external eligible-population denominator."),
    ("Branches", "No", "branches", "Passed through as a reported branch count.", "Measures reported branch-network scale; no branch addresses are available."),
    ("Employees (FTE)", "No", "employees", "Passed through as reported employee/FTE count.", "Provides institution staffing scale; definition follows the source extract."),
    ("Total loans ($M)", "No", "loansM", "Passed through as total loans in $ millions.", "Measures loan portfolio scale; product/category balances are not provided."),
    ("Deposits / shares ($M)", "No", "depositsM", "Passed through as deposits/shares in $ millions.", "Measures member funding scale; source combines deposits/shares as defined in the extract."),
    ("Net income ($M)", "No", "niM", "Passed through as period net income in $ millions.", "Shows current-period earnings; use with period and balance-sheet ratios."),
    ("Equity / net worth ($M)", "No", "equityM", "Passed through unless separate equity amount/earnings/reserve/interest/gain/loss/merger components sum to nonzero and source equity is absent or >= 50% of assets; then the component sum is substituted.", "Capital dollars provide loss-absorption context. Component substitution is an import normalization rule."),
    ("Return on assets (%)", "No", "roa; fallback: niM, assetsM", "Uses source roa when present; otherwise niM / assetsM x 100.", "Profitability relative to asset scale. Fallback uses point-in-time assets, not average assets."),
    ("Return on equity (%)", "No", "roe; fallback: niM, equityM", "Uses source roe when present; otherwise niM / equityM x 100.", "Earnings relative to reported net worth/equity."),
    ("Loan-to-share ratio (%)", "No", "ltd; fallback: loansM, depositsM", "Uses source ltd when present; otherwise loansM / depositsM x 100.", "Shows loan funding relative to member shares/deposits; interpret with liquidity and funding strategy."),
    ("Loan-to-asset ratio (%)", "No", "lta; fallback: loansM, assetsM", "Uses source lta when present; otherwise loansM / assetsM x 100.", "Shows the share of assets represented by loans; higher deployment can reduce liquidity flexibility."),
    ("Delinquency rate (%)", "No", "dr; fallback: delinquentAmountM, loansM", "Uses source dr when present; otherwise delinquentAmountM / loansM x 100.", "Credit-quality indicator for delinquent balances as a share of loans."),
    ("Net charge-off rate (%)", "No", "nco; fallback: chargeOffM, recoveryM, loansM", "Uses source nco when present; otherwise (chargeOffM - recoveryM) / loansM x 100.", "Net realized loan-loss rate relative to loans; period convention follows source."),
    ("Portfolio yield (%)", "No", "py; fallback: grossIncomeM, assetsM", "Uses source py when present; otherwise grossIncomeM / assetsM x 100.", "Fallback is an income-to-assets yield proxy, not product-level loan yield."),
    ("Net worth / capital ratio (%)", "No", "capRatio; fallback: equityM, assetsM", "Uses source capRatio when present; otherwise equityM / assetsM x 100.", "Capital adequacy ratio used for resilience benchmarking."),
    ("Net interest income ($M)", "No", "niiM", "Passed through as reported net interest income in $ millions.", "Core spread-income measure; not a standalone Competitive Intelligence chart."),
    ("CEO", "No", "ceo", "Passed through as institution leadership text.", "Institution profile metadata; not a competitive KPI."),
    ("Website", "No", "website", "Passed through as URL/text.", "Institution profile/digital-presence reference; not a digital engagement measure."),
    ("Gross income ($M)", "No", "grossIncomeM", "Passed through as reported gross income in $ millions.", "Denominator for efficiency and net-margin calculations."),
    ("Operating expense ($M)", "No", "operatingExpenseM", "Passed through as operating expense in $ millions.", "Numerator for dashboard efficiency ratio."),
    ("Delinquent loan amount ($M)", "No", "delinquentAmountM", "Passed through as delinquent balances in $ millions.", "Absolute troubled-balance exposure; interpret alongside delinquency rate and portfolio size."),
    ("Delinquent loan count", "No", "delinquentCount", "Passed through as reported number of delinquent loans.", "Count-based exposure context; account definition follows source."),
    ("Charge-offs ($M)", "No", "chargeOffM", "Passed through as reported charge-offs in $ millions.", "Gross realized loan losses for the source period."),
    ("Recoveries ($M)", "No", "recoveryM", "Passed through as reported recoveries in $ millions.", "Previously charged-off balances recovered in the source period."),
    ("Efficiency ratio (%)", "Yes", "operatingExpenseM, grossIncomeM", "Recomputed when both inputs are mapped: grossIncomeM > 0 ? operatingExpenseM / grossIncomeM x 100 : 0. Overrides any supplied efficiencyRatio.", "Operating expense per dollar of gross income; lower is generally more efficient. Validate definitions before external reporting."),
    ("Net income margin (%)", "Yes", "niM, grossIncomeM", "Recomputed when both inputs are mapped: grossIncomeM != 0 ? niM / grossIncomeM x 100 : 0. Overrides any supplied netMargin.", "Net income retained per dollar of gross income; negative values indicate a net loss relative to income."),
    ("Recovery rate (%)", "Yes", "recoveryM, chargeOffM", "Recomputed when both inputs are mapped: chargeOffM > 0 ? recoveryM / chargeOffM x 100 : 0. Overrides any supplied recoveryRate.", "Share of charge-offs offset by recoveries; validate period alignment and source definitions."),
    ("Charter / CU number", "No", "charter", "Passed through; preferred within-quarter de-duplication key when populated.", "Institution identifier supporting stable joins across reporting quarters."),
    ("Reporting quarter", "No", "quarter", "Source period label; current workbook contains only 2026 Q1.", "Defines the observation period and is required for longitudinal comparisons."),
]

DERIVED_METRICS = [
    ("Amount unit normalization", "Yes", "assetsM and MONEY_KEYS fields", "For imported data, if median positive assets > 10,000,000 divide money fields by 1,000,000; else if > 5,000 divide by 1,000; otherwise leave unchanged.", "Converts dollars or thousands into $ millions; assumes comparable units across listed money fields."),
    ("Ratio unit normalization", "Yes", "ltd, lta, dr, nco, py, capRatio, roa, roe", "If all available median tests (ltd < 3, capRatio < 0.5) are true, multiply ratio keys by 100.", "Normalizes imported fractional ratios to percentage points; assumes ratio fields use a consistent scale."),
    ("Peer comparison universe", "Yes", "bucket, state, selected CU", "Standard peer tabs use selected bucket, optionally adjacent buckets and same-state filter. Competitive scopes: same state/all buckets; nationwide buckets within +/-1; or nationwide. Peer statistics exclude selected CU.", "Makes benchmark population explicit. Competitive market shares/ranks include the CU in the market universe."),
    ("Universe institution count", "Yes", "name, state, bucket", "Count rows after applying the selected scope.", "Context for market scale and statistical reliability."),
    ("Universe total assets ($M)", "Yes", "assetsM", "Sum assetsM for all institutions in selected competitive universe, including selected CU.", "Denominator for selected-universe asset share."),
    ("Universe total members", "Yes", "members", "Sum members for all institutions in selected competitive universe, including selected CU.", "Denominator for member share; not an addressable population."),
    ("Competitive asset rank", "Yes", "assetsM, selected scope", "1 + count(universe assetsM > selected CU assetsM); ties share competition rank.", "Shows asset scale position in the displayed market/peer universe."),
    ("Competitive asset percentile", "Yes", "assetsM, selected scope", "count(universe assetsM <= selected CU assetsM) / universe count x 100; includes selected CU.", "Indicates relative asset scale; higher means more institutions are at or below the selected CU."),
    ("Asset share of universe (%)", "Yes", "assetsM, selected scope", "selected CU assetsM / sum(universe assetsM) x 100. Labeled in-state share only in state scope.", "Measures reported CU asset concentration in the selected universe, not all financial institutions."),
    ("Member share of universe (%)", "Yes", "members, selected scope", "selected CU members / sum(universe members) x 100.", "Membership concentration among included CUs, not community penetration."),
    ("Closest asset-size rivals", "Yes", "assetsM, selected scope", "Sort peers by abs(peer assetsM - selected CU assetsM); take first three when no named rivals are selected.", "Size-proximate shortlist; does not imply geographic or business-line similarity."),
    ("Named-rival comparison table", "Yes", "assetsM, members, efficiencyRatio, capRatio, roa, dr, ltd", "Show selected CU, up to three selected rivals (or closest-size defaults), and field-wise peer median.", "Lets leaders inspect named institutions against the broader benchmark."),
    ("Peer median", "Yes", "peer values for displayed metric", "Sort ascending and select index floor(0.5 x (N - 1)); if N=0, helper returns 0.", "Robust central benchmark. Empty-group zero is only a code fallback and means no benchmark."),
    ("Peer 25th / 75th percentile", "Yes", "peer values for displayed metric", "Sort ascending and select index floor(p x (N - 1)) for p=0.25 or 0.75; discrete, not interpolated.", "Shows peer distribution bounds used by board, risk, and opportunity views."),
    ("Metric percentile rank", "Yes", "peer values, selected CU metric", "count(peer values <= selected CU value) / peer count x 100. Display is inverted as 100 - percentile for dr, nco, ltd.", "Places CU within peer distribution while adjusting for lower-is-better ratios."),
    ("Efficiency ratio gap (pp)", "Yes", "efficiencyRatio, peer efficiencyRatio", "selected CU efficiencyRatio - peer median; positive indicates higher operating expense per income dollar.", "Quantifies operating-efficiency distance; lower is generally better."),
    ("Loans-to-assets ratio (%)", "Yes", "loansM, assetsM", "loansM / assetsM x 100 when assetsM > 0; otherwise 0 in chart. Peer reference is median of institution-level ratios.", "Indicates lending deployment; not a loan product mix."),
    ("Deposits-to-assets ratio (%)", "Yes", "depositsM, assetsM", "depositsM / assetsM x 100 when assetsM > 0; otherwise 0. Peer reference is median of institution-level ratios.", "Shows member funding relative to assets; independent ratio, not an additive asset component."),
    ("Competitive net-worth benchmark", "Yes", "capRatio", "Selected CU source capRatio versus discrete median peer capRatio.", "Compares capital resilience within the selected universe."),
    ("Capital review signal", "Yes", "capRatio", "Signal when selected CU capRatio < 7%; otherwise above the 7% reference.", "Prompts review against dashboard's well-capitalized reference; not a regulatory determination."),
    ("Efficiency review signal", "Yes", "efficiencyRatio, peer efficiencyRatio", "Signal when CU efficiencyRatio > peer median.", "Flags relative operating-cost pressure; does not adjust for business model or accounting differences."),
    ("Delinquency review signal", "Yes", "dr, peer dr", "Signal when CU delinquency rate > peer discrete 75th percentile.", "Highlights a relative credit-quality outlier for review."),
    ("Top institutions by assets chart", "Yes", "name, assetsM, selected CU index", "Sort universe by assets descending; show top six, or top five plus selected CU if outside top six.", "Shows market scale leaders while keeping selected CU visible."),
    ("Balance-sheet structure chart", "Yes", "loansM, depositsM, assetsM, capRatio", "Selected CU and peer medians for loans/assets, deposits/assets, net-worth/assets. Independent ratios, not a 100% stacked composition.", "Compares lending deployment, member funding, and capital intensity; no product mix is implied."),
    ("State asset/member totals", "Yes", "state, assetsM, members, branches", "Group by state; count institutions and sum assets, members, branches.", "Describes headquarters-state concentration; branch totals are not geocoded."),
    ("State asset rank", "Yes", "state, assetsM", "Position after sorting CUs in selected headquarters state by assets descending; ties follow stable row order.", "Shows asset scale position within the headquarters state."),
    ("State member rank", "Yes", "state, members", "Position after sorting CUs in selected headquarters state by members descending; ties follow stable row order.", "Shows membership scale position within the headquarters state."),
    ("State map share (%)", "Yes", "state aggregate, selected map metric", "Selected-state total / same metric summed across configured map tiles x 100. Metrics: CU count, members, assets, branches.", "Places headquarters-state scale in the map's configured geography."),
    ("State map rank", "Yes", "state aggregates, selected map metric", "1 + count(configured state tile values > selected state value).", "Ranks state totals for the selected map metric."),
    ("Members per branch", "Yes", "members, branches", "members / branches when branches > 0. Peer reference uses median members / median branches.", "Approximate service-network intensity; not branch capacity or demand."),
    ("Members per employee", "Yes", "members, employees", "members / employees when employees > 0. Peer reference uses median members / median employees.", "Approximate staffing scale/productivity context; workforce definition follows source."),
    ("Member scale vs peer median (%)", "Yes", "members, peer median members", "(selected CU members / peer median members - 1) x 100 when median is nonzero.", "Shows relative membership scale to the selected peer group."),
    ("Branch scale vs peer median (%)", "Yes", "branches, peer median branches", "(selected CU branches / peer median branches - 1) x 100 when median is nonzero.", "Shows relative reported branch count to peers."),
    ("Delinquent share of loans (%)", "Yes", "delinquentAmountM, loansM", "delinquentAmountM / loansM x 100 when loansM > 0; otherwise unavailable.", "Scales delinquent dollar exposure to portfolio size."),
    ("Gap vs peer median (pp)", "Yes", "roa, roe, capRatio, lta, ltd, dr, nco, py", "For each gap metric: selected value - discrete peer median.", "Expresses distance from peer median in percentage points."),
    ("Gap strength/opportunity classification", "Yes", "gap, metric direction", "Strength if higher-is-better gap >= 0 or lower-is-better gap <= 0. Higher: roa, roe, capRatio, lta, py. Lower: ltd, dr, nco.", "Prioritizes relative performance under existing rules; lower-is-better loan-to-share is a model assumption, not universal guidance."),
    ("Opportunity count", "Yes", "gap classifications", "Count non-null metrics classified as opportunity.", "Counts unfavorable gaps under configured directional rules."),
    ("Strength count", "Yes", "gap classifications", "Count non-null metrics classified as strength.", "Counts metrics at or better than peer median."),
    ("Largest opportunity magnitude", "Yes", "opportunity gaps", "Maximum absolute opportunity gap, displayed in percentage points.", "Surfaces largest relative gap; not a dollar-impact estimate."),
    ("Delinquency risk flag", "Yes", "dr, peer dr", "Raised when dr > max(peer P75, 1.5%).", "Flags elevated past-due performance against peer distribution and fixed floor."),
    ("Net charge-off risk flag", "Yes", "nco, peer nco", "Raised when nco > max(peer P75, 0.5%).", "Flags elevated realized losses against peers and fixed floor."),
    ("Capital adequacy risk flag", "Yes", "capRatio", "Critical <2%; high 2-<4%; high 4-<6%; medium 6-<7%; no flag >=7%.", "Rule-based capital screen; not an official PCA determination."),
    ("Earnings risk flag", "Yes", "roa", "High when ROA <0%; medium when 0%<=ROA<0.2%; otherwise none.", "Highlights losses or thin asset returns."),
    ("Liquidity/funding risk flags", "Yes", "ltd, lta", "Loan-to-share >100%: loans exceed shares; loan-to-assets >92%: limited flexibility; loan-to-share <35%: excess liquidity/low loan demand.", "Prompts review of liquidity/funding posture; interpretation depends on business model."),
    ("Active risk-flag count", "Yes", "capRatio, roa, dr, nco, ltd, lta, peer distributions", "Count all rules triggered by computeRiskFlags for the selected CU.", "Summary count of rule-based screening alerts."),
    ("Peer active-flag rate (%)", "Yes", "peer risk flags, peer count", "count(peers with >=1 active risk flag) / peer count x 100; zero if peer group empty.", "Shows prevalence of dashboard flags in peer cohort."),
    ("Quarterly metric series", "Yes", "quarter, selected metric, institution identity", "Match institution by case-insensitive name and state per quarter; calculate same-quarter peer median for selected metric.", "Supports trends only after multiple periods are loaded; source currently has one quarter."),
    ("Quarterly gap to peer median", "Yes", "quarterly metric, quarterly peer values", "Institution value - same-quarter peer median. No growth rate is calculated.", "Shows relative period distance; not a substitute for growth."),
    ("Growth rates and gainers/decliners", "No", "Requires comparable prior/current periods", "Not calculated: all 4,336 rows are 2026 Q1. With history, growth could be (current/prior - 1) x 100 with zero/missing handling.", "Unavailable until multiple comparable quarters exist; cross-sectional differences are not growth."),
    ("Loan product composition", "No", "Requires category-level loan balances", "Not calculated: only total loansM is supplied.", "Requires product balances such as real estate, auto, cards, and other loans."),
    ("Branch/service-area overlap", "No", "Requires branch addresses/geocodes and service-area/field-of-membership data", "Not calculated: city is blank; no branch locations are present.", "Required for geographic footprint and market-overlap analysis."),
]

VISUALS = [
    ("Competitive Intelligence: scope selector", "state, bucket", "State market = same state/all sizes; similar-size = nationwide buckets +/-1; nationwide = all. Market includes selected CU; peer medians exclude it.", "CompetitiveTab"),
    ("Competitive Intelligence: asset rank/percentile", "assetsM", "Rank = 1 + count strictly greater; percentile = count <= CU / universe N x 100.", "CompetitiveTab"),
    ("Competitive Intelligence: asset/member share KPIs", "assetsM, members", "Selected CU / sum across selected universe x 100.", "CompetitiveTab"),
    ("Competitive Intelligence: efficiency gap", "operatingExpenseM, grossIncomeM, efficiencyRatio", "Efficiency ratio = operating expense/gross income x 100; gap = CU - peer median.", "CompetitiveTab / normalizeQuarterRows"),
    ("Competitive Intelligence: market-scale chart", "name, assetsM", "Top six by universe assets; include selected CU if outside top six.", "CompetitiveTab"),
    ("Competitive Intelligence: balance-sheet ratios", "loansM, depositsM, assetsM, capRatio", "Loans/assets, deposits/assets, net-worth/assets; CU vs peer medians. Independent ratios, not stacked.", "CompetitiveTab"),
    ("Competitive Intelligence: rival scorecard", "name, state, assetsM, members, efficiencyRatio, capRatio, roa, dr, ltd", "Up to three named rivals or closest-size defaults; current values vs peer medians.", "CompetitiveTab"),
    ("Competitive Intelligence: risk/opportunity signals", "capRatio, efficiencyRatio, dr, peer distributions", "Below 7% capital, above-median efficiency, or delinquency above peer P75.", "Screening aid, not supervisory determination."),
    ("Overview profile and KPIs", "assetsM, members, branches, employees, niM, roa, capRatio, dr, ltd", "Direct values, peer medians, thresholds; assets/member = assetsM x 1,000,000 / members.", "OverviewTab"),
    ("Overview operating analytics", "efficiencyRatio, netMargin, delinquentAmountM, loansM, recoveryRate, chargeOffM, recoveryM, lta", "Derived operating ratios and delinquent amount / loans x 100.", "NewDataAnalytics"),
    ("Overview peer percentile bars", "roa, capRatio, dr, ltd, peer ratios", "Peer percentile; invert display for dr, nco, ltd.", "OverviewTab / PercentileBar"),
    ("Peer Benchmarking radar and delta", "roa, roe, capRatio, dr, nco, ltd, lta, py", "Adjusted percentile profile and CU-minus-peer-median comparisons.", "BenchmarkTab"),
    ("Peer group detail", "name, state, assetsM, members, roa, capRatio, dr, ltd", "Sortable peer table, max 25 records including selected CU.", "PeerTable"),
    ("Compare Institutions", "assetsM, members, branches, roa, roe, capRatio, dr, nco, ltd, lta, py, niM", "Direct current-period values for up to four institutions; assets chart.", "CompareTab"),
    ("Trend Comparison", "quarter and selected TREND_METRICS field", "Match institution by name/state and calculate quarter-specific peer median.", "TrendComparisonTab; no growth KPI currently calculated."),
    ("Demographic Views and map", "state, region, members, branches, employees, assetsM", "State aggregates/ranks/shares; member percentile, peer scale, members per branch/employee.", "DemographicViewsTab / StateMap"),
    ("Gap Analysis", "roa, roe, capRatio, lta, ltd, dr, nco, py", "CU-minus-peer-median gaps, direction-based classification, sorted opportunity list.", "GapAnalysisTab"),
    ("Risk Alerts", "capRatio, roa, dr, nco, ltd, lta", "Fixed thresholds and peer P75 rules; fraction of peers with >=1 flag.", "RiskTab / computeRiskFlags"),
    ("Board Report", "assetsM, members, niM, branches, METRIC_META fields", "CU values, peer medians, and peer P25-P75.", "BoardReportTab"),
    ("AI Recommendations context", "CU scale, METRIC_META fields, peer distributions, flags", "Prompt supplies CU ratios, peer median/P25/P75, scope, and flags; generated prose is model output, not KPI.", "AITab / LLM provider"),
]


def style_table(sheet, widths, row_height):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="C9A227"))
    sheet.row_dimensions[1].height = 34
    for row_idx in range(2, sheet.max_row + 1):
        for cell in sheet[row_idx]:
            cell.font = Font(name="Aptos", size=10, color="243242")
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if row_idx % 2 == 0 else "F4F7FA")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="D7DEE7"))
        sheet.row_dimensions[row_idx].height = row_height
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width


workbook = Workbook()
dictionary = workbook.active
dictionary.title = "KPI Dictionary"
dictionary.append(["Metric Name", "Derived in Dashboard? (Yes/No)", "Source Column(s)", "Calculation Logic / Transformations", "Business Explanation"])
for row in SOURCE_FIELDS + DERIVED_METRICS:
    dictionary.append(row)
style_table(dictionary, {"A": 37, "B": 29, "C": 39, "D": 82, "E": 70}, 58)

mapping = workbook.create_sheet("Visual Mapping")
mapping.append(["Dashboard Visual / Section", "Underlying Data Fields", "Calculation / Data Mapping", "Component / Notes"])
for row in VISUALS:
    mapping.append(row)
style_table(mapping, {"A": 46, "B": 58, "C": 88, "D": 44}, 52)

notes = workbook.create_sheet("Read Me")
for row in [
    ["CU360 KPI Dictionary"],
    ["Source workbook", "backend/data/cu360_data.xlsx · sheet CU360 data"],
    ["Snapshot profile", "4,336 institution records; 35 columns; 54 headquarters state/territory codes; all rows report 2026 Q1."],
    ["Derived flag convention", "Yes = computed, transformed, aggregated, ranked, classified, or conditionally derived by dashboard logic. No = source field displayed directly. Overrides/fallbacks are noted in calculation logic."],
    ["Units", "Monetary values are $ millions; ratios display as percentages or percentage points."],
    ["Peer statistics", "Values are sorted and indexed at floor(p x (N - 1)); percentile rank counts values <= selected CU. Quantiles are discrete, not interpolated."],
    ["Competitive scope", "State = same headquarters state; similar-size = nationwide asset buckets within +/-1; nationwide = all institutions. Market shares/ranks include selected CU; peer medians exclude it."],
    ["Data limitations", "One reporting quarter means no growth, longitudinal variance, or gainers/decliners. City is blank; there are no branch geocodes or loan-product balances."],
    ["Interpretation", "Rule-based risk/opportunity flags are management screening aids, not official regulatory determinations. Peer differences are descriptive, not causal."],
    ["Coverage", "Dictionary includes all source fields, dashboard calculations, peer/statistical transforms, risk/gap rules, and requested measures unavailable from this snapshot. Visual Mapping inventories dashboard sections and fields."],
]:
    notes.append(row)
notes.sheet_view.showGridLines = False
notes.column_dimensions["A"].width = 30
notes.column_dimensions["B"].width = 115
for row_idx in range(1, notes.max_row + 1):
    notes.row_dimensions[row_idx].height = 34
    for cell in notes[row_idx]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor="DCE8F2" if row_idx == 1 else ("F4F7FA" if row_idx % 2 == 0 else "FFFFFF"))
        cell.font = Font(name="Aptos", size=10, bold=(cell.column == 1), color="17324D" if cell.column == 1 else "243242")
notes["A1"].font = Font(name="Aptos Display", size=15, bold=True, color="17324D")

workbook.save(OUTPUT)
check = load_workbook(OUTPUT, read_only=True, data_only=True)
print(f"Created: {OUTPUT}")
print(f"KPI dictionary entries: {check['KPI Dictionary'].max_row - 1}")
print(f"Visual mappings: {check['Visual Mapping'].max_row - 1}")
print(f"Source fields documented: {len(SOURCE_FIELDS)}")
print(f"Sheets: {', '.join(check.sheetnames)}")
check.close()
