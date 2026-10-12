"""Create the two-sheet source-to-KPI dictionary from Q1/Q2 NCUA CSV files."""

import csv
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "CU360_Source_to_KPI_Data_Dictionary.xlsx"
SOURCES = {
    "Q1 2026": Path.home() / "Downloads" / "AI Labs" / "NCUA" / "Data" / "Q1 2026.csv",
    "Q2 2026": Path.home() / "Downloads" / "AI Labs" / "NCUA" / "Data" / "Q2 2026.csv",
}
EQUITY_SOURCE_COLUMNS = (
    "Equity_Acquired_in_Merger; Equity_Undivided_Earnings_Amount; Equity_Other_Reserves_Amount; "
    "Equity_Appropriation_for_Non_Conforming_Investments; Equity_Noncontrolling_Interest_in_Consolidated_Subsidiaries; "
    "Equity_Accumulated_Unrealized_Net_Gains_Losses_on_Cash_Flow_Hedges; Equity_Other_Comprehensive_Income; "
    "Equity_Accumulated_Unrealized_Gains_Losses_on_Available_for_Sale_Debt_Securities; "
    "Equity_Accumulated_Unrealized_Losses_for_OTTI"
)
PEER_RATIO_SOURCE_COLUMNS = (
    "Net_Income; Total_Assets; " + EQUITY_SOURCE_COLUMNS + "; "
    "Total_Amount_of_Delinquent_Loans_Leases; Amount_of_Total_Loans_Leases; "
    "Total_Amount_of_All_Loans_Charged_Off_YTD; Total_Amount_of_All_Loans_Recovered_YTD; "
    "Total_Shares_Amount; Gross_Income"
)
TREND_SOURCE_COLUMNS = (
    "Quarter; CU_NUMBER; Total_Assets; Total_Members; Amount_of_Total_Loans_Leases; "
    "Total_Shares_Amount; " + PEER_RATIO_SOURCE_COLUMNS
)

# Rows: KPI, Direct/Derived, exact source column names, business formula,
# plain-English purpose, reference range, assumptions, current logic, recommended logic.
KPI_ROWS = [
    ("Reporting period", "Direct", "Quarter; CYCLE_DATE", "Use the row's Quarter; use CYCLE_DATE to validate the reporting date.", "Identifies which call-report snapshot the values belong to.", "Q1 2026 or Q2 2026", "Filter to one period before calculating any KPI or benchmark.", "The packaged API starts with Q1 only; imported periods are held in browser memory.", "Load and persist each quarter as an authoritative, separately selectable period; never label one period while showing another."),
    ("Credit union identity", "Direct", "CU_NUMBER; CU_Name; STATE", "Match across periods by CU_NUMBER. If unavailable, use CU_Name plus STATE.", "Identifies the credit union and supports safe quarter-to-quarter matching.", "One unique CU_NUMBER per period", "Names can change; CU_NUMBER is the preferred stable identifier.", "Quarter matching now prefers CU_NUMBER (mapped to charter), then case-insensitive name and state.", "Continue to use CU_NUMBER for joins, de-duplication and growth calculations; report unmatched records."),
    ("Total assets", "Direct", "Total_Assets", "Use the reported period-end asset balance; convert source dollars to millions only for display.", "Measures the credit union's balance-sheet size.", "Greater than 0 for a valid ratio denominator", "The CSV values are dollar amounts formatted with currency symbols and Indian-style comma grouping.", "Imported dollars are converted to $ millions using a file-wide median-asset heuristic.", "Preserve exact source dollars for calculations; format as dollars/millions only at display time."),
    ("Total members", "Direct", "Total_Members", "Use the reported member count for the selected quarter.", "Measures membership scale.", "Non-negative whole number", "Not population penetration without an external eligible-population denominator.", "Displayed directly after number parsing.", "Keep as a whole-number count and document any definition changes between periods."),
    ("Total loans and leases", "Direct", "Amount_of_Total_Loans_Leases", "Use the reported total loans and leases balance.", "Measures total lending exposure and portfolio scale.", "Non-negative balance; zero makes loan-based ratios unavailable", "Source is dollar balance, not a loan count.", "Mapped as total loans; amounts converted to $ millions for dashboard fields.", "Retain raw dollar amount and flag zero/missing values before deriving ratios."),
    ("Total shares / deposits", "Direct", "Total_Shares_Amount", "Use the reported shares balance as the funding balance.", "Measures member share/deposit funding.", "Non-negative balance; zero makes loan-to-share unavailable", "This is the balance column. Do not substitute Number_Total_Shares, which counts share accounts.", "Now mapped by exact amount-column alias; this fixes the prior 677,820.7% result.", "Use Total_Shares_Amount as the denominator for loan-to-share and growth."),
    ("Share account count", "Direct", "Number_Total_Shares", "Use as the number of share accounts; do not use as a dollar balance.", "Counts share accounts, not dollars of member deposits.", "Non-negative whole number", "Not interchangeable with Total_Shares_Amount.", "Available in the source but not a dashboard KPI.", "Use only for account-count analysis, never as the loan-to-share denominator."),
    ("Branch count", "Direct", "Branches", "Use reported branch count for the selected CU and period.", "Shows the reported size of the branch network.", "Non-negative whole number", "No branch addresses are provided, so this is not a geocoded footprint.", "Displayed directly in Overview and Demographic views.", "Keep the period filter applied; do not infer branch locations or market overlap."),
    ("Net income", "Direct", "Net_Income", "Use the reported net-income amount for the source period.", "Shows reported earnings for the reporting cycle.", "Can be positive or negative", "Validate whether this field is year-to-date before comparing Q1 to Q2.", "Displayed directly; used as the numerator for ROA/ROE fallbacks and margin.", "Quarterize YTD values for sequential-quarter comparisons; keep the raw source value available."),
    ("Gross income", "Direct", "Gross_Income", "Use the reported gross-income amount for the source period.", "Provides the income base used by efficiency and net-margin calculations.", "Non-negative for efficiency ratio; zero means unavailable", "Confirm the institution's approved efficiency-ratio denominator includes exactly this source definition.", "Displayed directly and used as a denominator for efficiency, margin and the portfolio-yield proxy.", "Document and validate the denominator definition before external benchmarking."),
    ("Operating expenses", "Direct", "Operating_Expenses", "Use the reported operating expense amount for the source period.", "Measures operating cost before calculating efficiency.", "Non-negative; period basis must match Gross_Income", "YTD/period alignment matters for ratios and period trends.", "Displayed directly and used as the efficiency-ratio numerator.", "Use the matching reporting-period expense basis and preserve N/A for missing inputs."),
    ("Net worth / capital ratio", "Derived", EQUITY_SOURCE_COLUMNS + "; Total_Assets; NET_WORTH_AMOUNT (source candidate, often invalid)", "Sum applicable reported equity/net-worth components, including signed gains and losses; divide by Total_Assets; multiply by 100.", "Shows the capital cushion available to absorb losses relative to assets.", "NCUA PCA reference: 7% well-capitalized; 6-<7% adequately capitalized; 4-<6% undercapitalized; 2-<4% significantly undercapitalized; <2% critically undercapitalized.", "NET_WORTH_AMOUNT equals Total_Assets for Affinity in both quarters and for 3,986 Q1 records; equity component values are used as the numerator when source net-worth is implausible. Component inclusion/sign conventions require provider confirmation.", "Components / Total_Assets x 100; importer can derive after assembling component fields. Affinity: Q1 7.5762%; Q2 7.6621%.", "Use the NCUA-defined net-worth numerator and denominator as confirmed by the data provider; do not trust NET_WORTH_AMOUNT where it duplicates assets. Preserve negative components and audit outliers."),
    ("Delinquency rate", "Derived", "Total_Amount_of_Delinquent_Loans_Leases; Amount_of_Total_Loans_Leases", "Total delinquent loan/lease dollars divided by total loans/leases, multiplied by 100.", "Shows the portion of the loan portfolio that is delinquent.", "Usually below 2%; >3% is an absolute review threshold requested for this dashboard.", "Zero/missing loans means N/A, not 0%. The reported 0.93% Affinity Q2 value is correct after rounding.", "Derived as delinquent amount / total loans x 100 when no ratio field is provided. Affinity Q1 0.7214%; Q2 0.9257%.", "Keep this formula; use N/A for zero denominator and separate the absolute >3% alert from peer-relative alerting."),
    ("Loan-to-share ratio", "Derived", "Amount_of_Total_Loans_Leases; Total_Shares_Amount", "Total loans and leases divided by total shares amount, multiplied by 100.", "Shows how much member share funding is deployed into loans.", "User guide: 50-110% typical; >120% aggressive funding monitor.", "Never use Number_Total_Shares. A zero share balance means N/A.", "Derived from total loan dollars / total share dollars x 100. Affinity Q1 101.3878%; Q2 102.6001%.", "Same formula with exact Total_Shares_Amount denominator; add distinct >100% monitor and >120% high-risk alert only if approved by policy."),
    ("Loan-to-assets ratio", "Derived", "Amount_of_Total_Loans_Leases; Total_Assets", "Total loans and leases divided by total assets, multiplied by 100.", "Shows the share of assets deployed into loans.", "No universal threshold; interpret with liquidity and business model.", "Zero assets means N/A. Do not describe this as product composition.", "Derived as loans / assets x 100 when the ratio column is absent.", "Use the same formula; benchmark against comparable peers and monitor liquidity context."),
    ("Return on assets (ROA)", "Derived", "Net_Income; Total_Assets", "Current dashboard proxy: Net_Income divided by ending Total_Assets, multiplied by 100. Recommended: annualized period net income divided by average assets, multiplied by 100.", "Measures earnings generated for each dollar of assets.", "Typical guide supplied: 0.25-1.50%; benchmark depends on period and accounting definitions.", "Source Net_Income appears YTD across Q1/Q2; the CSV has no average-asset column. Q1/Q2 values should not be compared as standalone quarter earnings without annualization/average balances.", "Q1 Affinity 0.0477%; Q2 0.1044% unannualized using ending assets.", "Annualize the period's earnings consistently and divide by average assets; obtain opening/average asset balances. Until available, label the current figure unannualized proxy."),
    ("Return on equity (ROE)", "Derived", "Net_Income; " + EQUITY_SOURCE_COLUMNS + "; NET_WORTH_AMOUNT (source candidate, often invalid)", "Current dashboard fallback: Net_Income divided by reported equity, multiplied by 100. Recommended: annualized net income divided by average equity, multiplied by 100.", "Measures earnings relative to member capital.", "No single universal range; compare only with matching periods and definitions.", "YTD net income and ending equity can distort quarterly ROE; average equity is not supplied.", "Uses source ROE if present; otherwise net income / equity x 100.", "Annualize earnings and use average equity; confirm the correct equity numerator with the source provider."),
    ("Efficiency ratio", "Derived", "Operating_Expenses; Gross_Income", "Operating expenses divided by gross income, multiplied by 100.", "Shows how much operating expense is used to produce each dollar of gross income; lower is generally more efficient.", "Context-dependent; compare similar business models and consistent definitions.", "Both inputs must cover the same period and use the institution's approved definition.", "Recomputed from source expense/income amounts when both mapped; Affinity Q1 57.07%; Q2 57.54%.", "Formula is mathematically valid if Gross_Income is the approved denominator. Document inclusions and period basis."),
    ("Net income margin", "Derived", "Net_Income; Gross_Income", "Net income divided by gross income, multiplied by 100.", "Shows the share of gross income remaining as net income.", "Positive is generally preferable; compare on a consistent YTD or quarter basis.", "Gross income of zero means N/A.", "Recomputed from Net_Income / Gross_Income x 100; Affinity Q1 3.27%; Q2 3.56%.", "Keep only after confirming both source fields have the same period basis; return N/A for zero income."),
    ("Delinquent loan balance", "Direct", "Total_Amount_of_Delinquent_Loans_Leases", "Use reported delinquent loan/lease dollars; display in millions after scaling.", "Shows the absolute dollars past due.", "Non-negative in normal records; inspect unusually large values against loans.", "Not a rate; compare with total loans to assess materiality.", "Displayed directly as Delinquent Loans dollars.", "Keep as a direct amount and display beside the delinquency ratio."),
    ("Delinquent loan count", "Direct", "Number_of_Delinquent_Loans_Leases", "Use the reported count of delinquent loans/leases.", "Shows the number of delinquent accounts.", "Non-negative whole number", "Count is not equivalent to delinquent balance.", "Displayed directly in operating analytics.", "Keep as a separate count KPI; do not infer average account balance without dividing amount by count."),
    ("Charge-offs, year to date", "Direct", "Total_Amount_of_All_Loans_Charged_Off_YTD", "Use reported YTD charge-off amount.", "Shows gross loan losses charged off so far in the reporting year.", "Non-negative YTD amount", "YTD values accumulate through the year; Q2 minus Q1 is the Q2 incremental amount only if definitions align.", "Displayed as charge-offs; used by recovery-rate/NCO calculations after mapping.", "Label explicitly as YTD or quarterize before QoQ trend comparison."),
    ("Recoveries, year to date", "Direct", "Total_Amount_of_All_Loans_Recovered_YTD", "Use reported YTD recovered amount.", "Shows collections recovered on previously charged-off loans.", "Non-negative YTD amount", "YTD period basis must match charge-offs.", "Displayed as recoveries.", "Label YTD; quarterize by subtracting prior YTD values before quarter comparisons."),
    ("Net charge-off rate", "Derived", "Total_Amount_of_All_Loans_Charged_Off_YTD; Total_Amount_of_All_Loans_Recovered_YTD; Amount_of_Total_Loans_Leases", "Current proxy: (YTD charge-offs minus YTD recoveries) divided by ending loans, multiplied by 100. Recommended: annualized period net charge-offs divided by average loans, multiplied by 100.", "Measures realized net credit losses relative to the loan portfolio.", "Peer/distribution-based; higher rates warrant review.", "YTD charge-offs and recoveries are not a single-quarter numerator; the CSV has no average-loan column.", "Uses (charge-offs - recoveries) / loans x 100 when source ratio missing.", "Quarterize the YTD amounts, annualize consistently, and divide by average loans before comparing quarters."),
    ("Recovery rate", "Derived", "Total_Amount_of_All_Loans_Recovered_YTD; Total_Amount_of_All_Loans_Charged_Off_YTD", "Recoveries divided by charge-offs, multiplied by 100.", "Shows what portion of charged-off dollars was recovered.", "0-100% is a common interpretation; values above 100% are possible due to timing differences.", "Both source values are YTD; not a standalone quarter collection rate.", "Recomputed as recoveries / charge-offs x 100 when both source fields mapped.", "Keep as a YTD recovery ratio with matching periods; use N/A when charge-offs are zero."),
    ("Portfolio yield (dashboard proxy)", "Derived", "Gross_Income; Total_Assets", "Current proxy: Gross_Income divided by Total_Assets, multiplied by 100.", "Current value is gross income relative to assets; it is not necessarily the yield earned on loans.", "No expected range until the metric is correctly defined.", "The CSV has category loan rates/balances, but duplicates and incomplete mapping prevent a validated weighted loan yield.", "Shown as Portfolio Yield (`py`) and can be derived from gross income / assets x 100.", "Rename to Gross Income / Assets Proxy, or replace with loan interest income / average loans using provider-validated fields."),
    ("Cost of funds", "Derived (not implemented)", "No direct cost-of-funds or interest/dividend expense column in either CSV; Total_Shares_Amount is available but average funding balance is not.", "Recommended: interest/dividend expense on member shares and borrowings divided by average interest-bearing funding, multiplied by 100.", "Measures the average rate paid to obtain funding.", "Compare with funding strategy and market rates; no universal band.", "No source column for funding interest/dividend expense or average funding balance exists in the reviewed CSVs.", "Not displayed or calculated in current dashboard.", "Do not populate from loan-category Interest_Rate columns; obtain funding expense and average balances."),
    ("Asset growth, Q1 to Q2", "Derived (not displayed as KPI)", "Total_Assets in Q1 2026; Total_Assets in Q2 2026; CU_NUMBER", "(Q2 assets minus Q1 assets) divided by Q1 assets, multiplied by 100, after matching CU_NUMBER.", "Shows the percentage change in assets over the quarter-to-quarter reporting snapshots.", "No universal positive/negative target; growth should be interpreted with membership, funding and earnings.", "Q1/Q2 represent two reporting snapshots; confirm period comparability and no restatement.", "Trend tab reports the absolute first-to-last change; no growth percentage KPI.", "Use the stated percentage-change formula, return N/A for zero/missing Q1 assets, and report unmatched CUs separately."),
    ("Loan growth, Q1 to Q2", "Derived (not displayed as KPI)", "Amount_of_Total_Loans_Leases in Q1 2026; Amount_of_Total_Loans_Leases in Q2 2026; CU_NUMBER", "(Q2 loans minus Q1 loans) divided by Q1 loans, multiplied by 100, after matching CU_NUMBER.", "Shows growth or contraction of the loan balance.", "Interpret with delinquency, funding and risk appetite.", "Period-end balance growth, not loan originations or production.", "No percentage growth KPI shown.", "Use same-dollar-balance formula; N/A for zero/missing prior loans."),
    ("Deposit/share growth, Q1 to Q2", "Derived (not displayed as KPI)", "Total_Shares_Amount in Q1 2026; Total_Shares_Amount in Q2 2026; CU_NUMBER", "(Q2 shares minus Q1 shares) divided by Q1 shares, multiplied by 100, after matching CU_NUMBER.", "Shows change in member funding balances.", "Interpret with loan growth, liquidity and funding mix.", "Balance growth is not new accounts; do not use Number_Total_Shares.", "No percentage growth KPI shown.", "Use Total_Shares_Amount for both periods; N/A for zero/missing prior balance."),
    ("Membership growth, Q1 to Q2", "Derived (not displayed as KPI)", "Total_Members in Q1 2026; Total_Members in Q2 2026; CU_NUMBER", "(Q2 members minus Q1 members) divided by Q1 members, multiplied by 100, after matching CU_NUMBER.", "Shows change in reported member count.", "No universal target; interpret with market and member engagement data.", "Member-count changes may reflect corrections or mergers as well as organic growth.", "No percentage growth KPI shown.", "Use stable CU_NUMBER matching; report new/closed institutions separately."),
    ("Market share of CU assets", "Derived", "Total_Assets; STATE; asset-size segment (derived from Total_Assets); CU_NUMBER", "Selected CU assets divided by total assets for the currently selected comparison universe, multiplied by 100.", "Measures the CU's share of assets among the included credit unions, not share of all financial institutions.", "0-100% of chosen universe", "State scope includes all CUs in the state; similar-size scope uses adjacent asset buckets; national scope uses all CUs.", "Competitive tab computes asset share from filtered ACTIVE_CUS and includes selected CU in denominator.", "Formula is coherent for the defined CU universe; display denominator and do not label as broad market share."),
    ("Competitive loan ranking", "Derived (not implemented)", "Amount_of_Total_Loans_Leases; STATE; Total_Assets; CU_NUMBER", "Rank credit unions by total loan and lease balance within the selected state, size-peer or national universe.", "Shows relative lending scale among competitors.", "Rank 1 is largest", "Current Competitive Intelligence ranks only by assets; peer universe must be period-specific.", "No loan ranking is currently displayed.", "Add a separately labeled loan-balance ranking; exclude unmatched/missing values and disclose the comparison universe."),
    ("Competitive deposit/share ranking", "Derived (not implemented)", "Total_Shares_Amount; STATE; Total_Assets; CU_NUMBER", "Rank credit unions by total shares/deposits balance within the selected universe.", "Shows relative member-funding scale.", "Rank 1 is largest", "Use dollar balance, never Number_Total_Shares.", "No deposit/share ranking is currently displayed.", "Add a separately labeled deposit/share ranking with period-specific rows and explicit denominator/scope."),
    ("Competitive membership ranking", "Derived (not implemented)", "Total_Members; STATE; Total_Assets; CU_NUMBER", "Rank credit unions by reported member count within the selected universe.", "Shows relative membership scale among competitors.", "Rank 1 is largest", "Member share is not community penetration.", "Competitive tab has asset rank and member share, but no membership rank.", "Add membership rank separately; retain the selected CU and tie rule."),
    ("Competitive growth ranking", "Derived (not implemented)", "Quarter; CU_NUMBER; Total_Assets; Amount_of_Total_Loans_Leases; Total_Shares_Amount; Total_Members", "For each measure, calculate Q2 minus Q1 and divide by Q1; rank dollar/count change and percentage change separately among matched CUs.", "Identifies growth leaders and decliners without mixing unlike units.", "Rank separately by balance/count and percent", "Only CUs with both periods can be ranked; zero/missing Q1 values are N/A.", "No growth ranking is currently displayed.", "Add period-matched rankings with a clear common-CU population and show unmatched institutions separately."),
    ("Share of membership in CU universe", "Derived", "Total_Members; STATE; asset-size segment (derived from Total_Assets); CU_NUMBER", "Selected CU members divided by sum of members in the selected CU universe, multiplied by 100.", "Shows membership concentration among included credit unions.", "0-100% of chosen CU universe", "Not community penetration; no eligible population denominator is supplied.", "Computed using current market rows including selected CU.", "Retain as CU-universe share and label clearly; use an external population for penetration."),
    ("Assets rank and percentile", "Derived", "Total_Assets; STATE; asset-size segment (derived from Total_Assets); CU_NUMBER", "Rank = 1 + number of universe CUs with greater assets. Percentile = number with assets at or below selected CU divided by universe count, multiplied by 100.", "Shows relative asset scale within the selected universe.", "Rank 1 is largest; percentile 0-100", "Ties share the rank; percentile counts ties at-or-below.", "Competitive tab calculates rank/percentile from the active universe.", "Keep the documented tie rule and active-period universe; name this asset rank, not a composite rank."),
    ("Peer median for financial ratios", "Derived", PEER_RATIO_SOURCE_COLUMNS + "; CU_NUMBER; STATE; Total_Assets", "For ROA, ROE, capital, delinquency, NCO, loan-to-share, loan-to-assets and the gross-income/assets proxy, sort peer values and select the current discrete middle position; exclude selected CU.", "Provides the middle reference value for a defined peer set.", "Metric-specific", "Peers are selected asset bucket or adjacent buckets, optionally same state; exclude selected CU.", "Median helper uses floor(0.5 x (n-1)); differs from conventional averaged middle values for even groups.", "Use a documented conventional median (average the two middle values for even groups), or label the existing discrete method."),
    ("Peer percentile ranking", "Derived", PEER_RATIO_SOURCE_COLUMNS + "; CU_NUMBER; STATE; Total_Assets", "Count peer values less than or equal to selected value divided by peer count, multiplied by 100. Invert display for delinquency, NCO and loan-to-share.", "Shows where a CU sits in its peer distribution.", "0-100th percentile", "Ties count at-or-below; lower-is-better display applies to delinquency, net charge-off and loan-to-share.", "Percentile function is tie-inclusive; displayed adjusted percentile uses 100 minus raw percentile for lower-is-better fields.", "Use consistent ranking direction and use the same adjusted percentile for bar width, color and label."),
    ("Peer median gap / variance", "Derived", PEER_RATIO_SOURCE_COLUMNS + "; CU_NUMBER; STATE; Total_Assets", "Selected CU value minus peer median. Percentage-point units for ratios; absolute balance units for monetary metrics.", "Shows how far the CU is from the peer midpoint.", "Zero means equal to median", "A ratio difference is percentage points, not percent growth or percent variance.", "Several tabs show a signed difference from peer median; no universal variance percentage KPI is calculated.", "Label ratio gaps in percentage points. If percentage variance is desired, use `(CU - median) / median x 100` and handle zero median."),
    ("Competitive asset-size shortlist", "Derived", "Total_Assets; CU_NUMBER", "Select up to three peers with the smallest absolute difference between their assets and the selected CU's assets.", "Provides a starting list of similarly sized competitors.", "Three peers when available", "It is size proximity only, not geographic or business-model equivalence.", "Uses absolute asset distance within the active competitive universe.", "Keep labeled as closest asset-size peers; allow explicit rival selection as provided."),
    ("Top-six asset scale chart", "Derived", "CU_Name; Total_Assets; STATE; asset-size segment (derived from Total_Assets); CU_NUMBER", "Sort selected market universe by assets descending; show top six and include selected CU if outside that set.", "Shows market size leaders alongside the selected credit union.", "Six bars", "Selected CU highlighted; chart uses period-end assets.", "Recomputes from the active competitive universe.", "Retain active-period scope and label the chosen market universe."),
    ("Loans-to-assets structure ratio", "Derived", "Amount_of_Total_Loans_Leases; Total_Assets", "Loans and leases divided by total assets, multiplied by 100.", "Compares lending deployment relative to asset size.", "No universal target", "Not product mix; zero assets yields N/A.", "Competitive chart derives this ratio for the CU and takes the median of peer-level ratios.", "Keep as an independent ratio and clearly distinguish peer median from a ratio of peer totals."),
    ("Deposits-to-assets structure ratio", "Derived", "Total_Shares_Amount; Total_Assets", "Total shares divided by total assets, multiplied by 100.", "Compares member funding relative to assets.", "No universal target", "This is an independent balance-sheet ratio, not an additive asset component.", "Competitive chart derives this ratio and compares CU to peer median.", "Keep as independent ratio; use Total_Shares_Amount, not account count."),
    ("Net-worth-to-assets structure ratio", "Derived", EQUITY_SOURCE_COLUMNS + "; NET_WORTH_AMOUNT (source candidate); Total_Assets", "Validated net-worth components divided by total assets, multiplied by 100.", "Compares capital intensity with peer institutions.", "NCUA capital bands apply to net worth ratio", "Must use provider-validated component composition.", "Competitive chart compares capRatio with peer median.", "Use same validated numerator as Net Worth Ratio KPI."),
    ("Competitive efficiency gap", "Derived", "Operating_Expenses; Gross_Income; peer CU Operating_Expenses; peer CU Gross_Income", "Selected CU efficiency ratio minus peer median efficiency ratio, in percentage points.", "Shows relative operating cost efficiency; a positive gap means higher cost ratio.", "Lower than peer median is generally favorable", "Does not adjust for business model or reporting definitions.", "Uses selected CU ratio minus peer median; warns when selected ratio is higher.", "Retain as percentage-point gap and compare equivalent cohorts."),
    ("Named-rival scorecard", "Derived", "CU_NUMBER; CU_Name; STATE; Total_Assets; Total_Members; Operating_Expenses; Gross_Income; NET_WORTH_AMOUNT; " + EQUITY_SOURCE_COLUMNS + "; Net_Income; Total_Amount_of_Delinquent_Loans_Leases; Amount_of_Total_Loans_Leases; Total_Shares_Amount", "Compare selected CU and up to three explicitly selected competitors for current-period values; display peer-universe medians alongside.", "Lets business users validate chosen competitors against a consistent set of scale, earnings, capital, credit-risk and funding metrics.", "Metric-specific", "Closest asset-size rivals are only a default shortlist; identity and active quarter matter.", "Competitive scorecard shows current values for named rivals and the active peer median.", "Keep all selected rivals in the same period and show source values/formulas for each ratio."),
    ("Balance-sheet growth comparison", "Derived (not displayed as KPI)", "Total_Assets; Amount_of_Total_Loans_Leases; Total_Shares_Amount; Total_Members; CU_NUMBER; Quarter", "For each balance/count: `(Q2 value - Q1 value) / Q1 value x 100`, matching CU_NUMBER.", "Shows quarter-over-quarter change in scale, lending, member funding or membership.", "No universal target", "Only 4,295 CUs are present in both files; Q1 has 41 only and Q2 has 4 only.", "Trend view currently shows a first-to-last absolute change, not a growth percentage; available periods are in-memory after import.", "Add percentage growth and absolute change as separate measures; N/A for missing/zero prior values; label Q1/Q2 clearly."),
    ("Top gainers / decliners ranking", "Derived (not displayed)", "CU_NUMBER; CU_Name; STATE; Total_Assets; Amount_of_Total_Loans_Leases; Total_Shares_Amount; Total_Members; Quarter", "For each metric, rank matched CUs by absolute change and percentage change between Q1 and Q2.", "Identifies institutions with largest expansion/contraction.", "Not applicable; rank separately by metric", "Do not mix asset dollars, member counts and percentages in one ranking. Exclude unmatched institutions or list separately.", "Not implemented.", "Show dollar/count change and percentage change separately, with consistent CU_NUMBER joins and period scope."),
    ("Asset per member", "Derived", "Total_Assets; Total_Members", "Total assets in dollars divided by reported member count.", "Shows balance-sheet scale per member.", "No universal target", "Do not divide dollars already expressed in $ millions unless converted back to dollars.", "Overview displays assetsM x 1,000,000 / members.", "Keep currency units explicit; N/A when member count is zero."),
    ("Members per branch", "Derived", "Total_Members; Branches", "Total members divided by branch count.", "Approximate member scale per reported branch.", "No universal target", "No branch locations/capacity data; zero branches means N/A.", "Demographic view uses selected CU members/branches; peer reference is ratio of medians.", "Keep as an operating-scale proxy, not branch productivity or geographic coverage."),
    ("Members per employee", "Derived (not available in these CSVs)", "Total_Members; employee/FTE source column not present in Q1/Q2 CSVs", "Members divided by FTE employees.", "Would approximate member load per employee.", "No universal target", "The source files reviewed do not contain employee counts.", "Normalized dashboard may have an employees field in a different template; current Q1/Q2 import reports it missing and displays 0.", "Show N/A until a validated employee field is supplied; do not interpret 0 as zero staff."),
    ("State totals and map", "Derived", "STATE; Total_Assets; Total_Members; Branches; CU_NUMBER", "For each headquarters state, count CU_NUMBER records and sum assets, members and reported branches for the active period.", "Compares institution and branch-reported scale across states.", "Not applicable", "Geography is headquarters state only; state-wide map intentionally aggregates all active-period CUs.", "State map rebuilds from ACTIVE_CUS and offers count/member/assets/branch metrics.", "Keep period-scoped, distinguish CU headquarters from branch locations, and disclose configured geography coverage."),
    ("State market rank/share", "Derived", "STATE; selected state totals of Total_Assets, Total_Members, Branches or CU_NUMBER", "Rank selected state by selected total; state share = its total divided by total across mapped state tiles, multiplied by 100.", "Shows relative headquarters-state concentration for the selected map metric.", "Rank among mapped states; share 0-100% of mapped tiles", "Map denominator includes configured US state/territory tiles only.", "Ranks state totals and reports their share of mapped total.", "Continue to apply active quarter and clearly disclose the mapped-state universe."),
    ("Member percentile and scale vs peers", "Derived", "Total_Members; STATE; asset-size segment (derived from Total_Assets); CU_NUMBER", "Percentile = peers with members <= CU members / peer count x 100. Relative scale = (CU members / peer median - 1) x 100.", "Shows membership size relative to the selected peer set.", "Percentile 0-100", "Selected CU excluded from peer values; missing/zero median means N/A.", "Demographics uses peer median and percentile functions.", "Use the active quarter and same peer definition shown to users."),
    ("Asset rank within state", "Derived", "STATE; Total_Assets; CU_NUMBER", "Rank selected CU among same-state CUs by descending assets.", "Shows relative balance-sheet scale in the headquarters state.", "Rank 1 is largest", "Tie handling should be explicit; current position may follow row order for ties.", "Demographic/competitive tab calculate state ranks from active state rows.", "Use competition rank with ties sharing rank and stable CU identity."),
    ("Member rank within state", "Derived", "STATE; Total_Members; CU_NUMBER", "Rank selected CU among same-state CUs by descending members.", "Shows relative membership scale in the headquarters state.", "Rank 1 is largest", "Tie handling should be explicit.", "Demographic view ranks active state rows.", "Use tie-aware rank and period-filtered dataset."),
    ("Gap-analysis metric variance", "Derived", PEER_RATIO_SOURCE_COLUMNS, "For each gap metric, subtract peer median from CU value. Current metrics include ROA, ROE, net worth ratio, loan-to-assets, loan-to-share, delinquency, NCO and gross-income/assets proxy.", "Highlights relative strengths and unfavorable gaps.", "Gap is measured in percentage points for ratios.", "Directional assumptions: higher is better for ROA/ROE/capital/LTA/proxy; lower for loan-to-share/delinquency/NCO. Lower loan-to-share is not universally better at every level.", "Uses fixed higher/lower direction per metric and classifies opportunity/strength.", "Keep labels as heuristic, review the direction for loan-to-share and replace the proxy with a validated measure."),
    ("Opportunity and strength counts", "Derived", PEER_RATIO_SOURCE_COLUMNS, "Count each non-null gap classified unfavorable/favorable for ROA, ROE, net worth ratio, loan-to-assets, loan-to-share, delinquency, NCO and gross-income/assets proxy.", "Summarizes how many metrics lag or meet/exceed peer median.", "0 through number of gap metrics", "Counts metrics, not dollar impact or materiality.", "Counts classified gap rows.", "Do not classify a missing value as an opportunity or strength; use N/A."),
    ("Largest opportunity magnitude", "Derived", PEER_RATIO_SOURCE_COLUMNS, "Largest absolute unfavorable CU-minus-peer-median gap, displayed in percentage points.", "Surfaces the largest relative ratio gap for review.", "Not a dollar impact estimate", "All included metrics are ratios; gap magnitude is not weighted by financial materiality.", "Sorts opportunities by absolute gap.", "Label as largest ratio gap and consider materiality-weighted prioritization separately."),
    ("Capital alert / health class", "Derived", EQUITY_SOURCE_COLUMNS + "; NET_WORTH_AMOUNT (source candidate); Total_Assets", "Current dashboard: <2% critical; 2-<4% high; 4-<6% high; 6-<7% medium; >=7% no capital flag.", "Classifies capital adequacy and highlights low net worth ratio.", "NCUA PCA thresholds: 7%, 6%, 4%, 2% bands as documented above.", "Requires valid net-worth components and assets. User-proposed <5% critical threshold differs from PCA bands.", "Capital flag uses the existing threshold tiers; competitive card uses <7% risk condition.", "Use official threshold bands with institution-specific regulatory applicability; avoid one generic 'critical <5%' label."),
    ("Earnings alert", "Derived", "Net_Income; Total_Assets", "Current: high when ROA <0%; medium when 0%<=ROA<0.2%; otherwise none.", "Highlights losses or thin returns relative to assets.", "Dashboard threshold 0% / 0.2%; user guidance says typical ROA 0.25-1.50%.", "Current ROA is not annualized and uses ending assets; average assets are not in the CSV.", "Rule-based on current ROA.", "Apply thresholds only to a validated annualized ROA basis and distinguish management alert from regulatory category."),
    ("Delinquency alert", "Derived", "Total_Amount_of_Delinquent_Loans_Leases; Amount_of_Total_Loans_Leases", "Current high alert when CU delinquency > max(peer 75th percentile, 1.5%). Recommended: peer outlier alert plus separate absolute alert above 3%.", "Identifies elevated past-due credit exposure.", "Typically <2%; >3% absolute review trigger per request.", "Zero loans means N/A; peer 75th percentile is derived from the same two source fields across peer CUs. Current formula has no unconditional >3% trigger.", "Peer P75/1.5% rule; Competitive Intelligence separately signals above peer P75.", "Implement both peer-relative and absolute guardrails and suppress if source denominator is invalid."),
    ("Net charge-off alert", "Derived", "Total_Amount_of_All_Loans_Charged_Off_YTD; Total_Amount_of_All_Loans_Recovered_YTD; Amount_of_Total_Loans_Leases", "Current alert when CU NCO > max(peer P75, 0.5%).", "Identifies elevated realized credit losses.", "Peer-relative threshold; no universal range specified.", "YTD numerator over ending balance makes quarter alerts non-comparable; peer P75 is derived from the same YTD fields for peer CUs.", "Uses peer percentile plus fixed floor.", "Quarterize/annualize net charge-offs and use average loans before alerting."),
    ("Loan-to-share / liquidity alerts", "Derived", "Amount_of_Total_Loans_Leases; Total_Shares_Amount; Total_Assets", "Current: loan-to-share >100% medium; <35% low; loan-to-assets >92% low.", "Flags high loan funding reliance or potentially under-deployed member funding.", "User guide: 50-110% typical; >120% aggressive.", "Thresholds are management screens, not regulatory standards; direction is not monotonic at all values.", "computeRiskFlags applies the stated >100/<35 LTD and >92 LTA rules.", "Use staged ranges such as monitoring >100%, heightened >110%, aggressive >120%, calibrated to policy and peer context."),
    ("Active risk flags / highest severity", "Derived", PEER_RATIO_SOURCE_COLUMNS, "Count triggered rules; highest severity is the top-ranked active flag (Critical, High, Medium, Low).", "Summarizes the number and highest priority of screening indicators.", "Zero flags means no configured rule fired, not that no risk exists.", "Depends on each component formula and threshold; missing values must not be treated as zero.", "Count and severity from rule-based thresholds on the listed source fields.", "Keep severity ordering, expose rule names/thresholds, and distinguish missing data from safe status."),
    ("Peer active-risk rate", "Derived", PEER_RATIO_SOURCE_COLUMNS + "; CU_NUMBER; STATE; Total_Assets", "Number of peers with at least one active risk flag divided by peer count, multiplied by 100.", "Shows prevalence of dashboard-defined flags among peers.", "0-100%", "Peer set excludes selected CU and follows current asset/state filters.", "Calculates share of peers with one or more configured flags.", "Retain same-period peer set and label as dashboard-rule flag rate, not regulatory distress rate."),
    ("Trend value by quarter", "Derived", TREND_SOURCE_COLUMNS, "Match CU by CU_NUMBER for each loaded period; display the selected period value and same-quarter peer median.", "Shows how a selected balance or ratio changes across loaded quarters.", "Metric-specific", "Only Q1/Q2 currently supplied; some income/loss measures are YTD.", "Trend series matches by stable CU identity after this audit; peer median recomputed per quarter.", "Keep snapshots separated by period and clearly identify YTD versus period-end measures."),
    ("Trend change over period", "Derived", TREND_SOURCE_COLUMNS, "Current dashboard: last available value minus first available value (absolute change). Recommended growth: (current - prior) / prior x 100.", "Shows the size and direction of change; a percentage change is easier to compare across institutions.", "Not applicable; depends on metric and baseline", "Zero or missing prior value means percentage growth N/A.", "KPI is an absolute delta, not a percentage variance.", "Show absolute delta and percent change in separate KPIs; match same CU and adjacent quarter."),
    ("Trend coverage", "Derived", TREND_SOURCE_COLUMNS, "Number of quarters with a matched CU value divided by number of registered quarters.", "Shows whether the selected CU has data for every loaded period.", "0/N to N/N quarters", "A CU absent in a quarter is missing, not zero.", "Counts non-null selected CU points over registered quarters.", "Keep missing quarters visibly missing; do not connect across gaps as if data existed."),
    ("Board report peer quartiles", "Derived", PEER_RATIO_SOURCE_COLUMNS + "; CU_NUMBER; STATE; Total_Assets", "Report selected CU value, peer discrete median, P25 and P75 for ROA, ROE, capital, delinquency, NCO, loan-to-share, loan-to-assets and portfolio-yield proxy.", "Provides a concise peer context for board review.", "P25-P75 is the central half of peer observations.", "Current helper uses discrete floor-index quantiles, not interpolation; peer group excludes selected CU.", "Board report uses selected quarter peers.", "Use one documented quantile method consistently and include period and peer-group definition."),
    ("Compare-tab best performer", "Derived", "Total_Assets; Total_Members; Branches; Net_Income; " + PEER_RATIO_SOURCE_COLUMNS + "; CU_NUMBER", "Current dashboard chooses max for most metrics and min only for delinquency, NCO and loan-to-share.", "Highlights the numeric best among selected CUs for each comparison row.", "Metric-specific", "Max is not automatically better for every metric; lower costs/risks and some ratios have optimum bands.", "Best performer classification uses a narrow lower-is-better set.", "Define direction per metric; use an acceptable range for liquidity ratios instead of a simplistic maximum."),
    ("AI recommendation output", "Derived narrative (not numeric KPI)", "CU_Name; STATE; Total_Assets; Total_Members; Branches; " + PEER_RATIO_SOURCE_COLUMNS + "; CU_NUMBER", "Model-generated recommendations are based on current CU context, peer statistics and active flags; no deterministic formula.", "Turns measured gaps and flags into suggested management actions.", "Not applicable", "Generated text must be validated against the supplied facts and approved business policy.", "Prompt passes current CU ratios, peer distributions and flags to the configured provider.", "Treat as advisory narrative, keep source figures visible, and do not present model output as calculated or regulatory fact."),
]

SOURCE_DESCRIPTIONS = {
    "Quarter": ("Reporting quarter label for this credit-union snapshot.", "Reporting period, growth and trend calculations."),
    "CU_NUMBER": ("Unique NCUA credit-union identifier.", "Institution matching, period joins, rankings and de-duplication."),
    "JOIN_NUMBER": ("Source-system join identifier; exact business definition should be confirmed with provider.", "Not currently used in dashboard KPIs."),
    "CYCLE_DATE": ("Call-report cycle/as-of date.", "Reporting-period validation and quarter identification."),
    "CU_Name": ("Reported credit-union name.", "Search, display labels, institution identity fallback, rival comparisons."),
    "STATE": ("Headquarters state code.", "State market filter, geographic aggregation, peer filtering."),
    "CEO": ("Reported chief executive officer name.", "Institution profile only."),
    "Website": ("Reported institution website URL.", "Institution profile link only."),
    "Branches": ("Reported number of branches.", "Overview profile, demographic KPI, members-per-branch, state branch aggregation."),
    "Total_Members": ("Reported total member count.", "Member KPI, member share, ranking, demographic map, membership growth."),
    "Total_Assets": ("Reported total assets in dollars.", "Asset size, size segment, ratios, rankings, market shares, growth."),
    "Amount_of_Total_Loans_Leases": ("Reported total loan and lease dollar balance.", "Loan-to-share, delinquency, loan-to-assets, net charge-off denominator, loan growth."),
    "Number_of_Total_Loans_Leases": ("Reported count of total loans and leases.", "Not currently used in dashboard KPIs."),
    "Number_Total_Shares": ("Number of member share accounts; this is a count, not a dollar balance.", "Not a loan-to-share denominator; no current KPI uses it."),
    "Total_Shares_Amount": ("Reported total share/deposit dollar balance.", "Loan-to-share denominator, deposits/assets, market share and deposit growth."),
    "NET_WORTH_AMOUNT": ("Reported net-worth amount; equals total assets in many rows reviewed and requires validation.", "Capital ratio source candidate; dashboard uses equity components when source amount appears implausible."),
    "Gross_Income": ("Reported gross income for the call-report period (confirm YTD definition).", "Efficiency ratio denominator, net income margin denominator, current portfolio-yield proxy."),
    "Net_Income": ("Reported net income for the call-report period (confirm YTD definition).", "Net income KPI; ROA/ROE and net margin numerators."),
    "Operating_Expenses": ("Reported operating expense for the call-report period.", "Efficiency ratio numerator."),
    "Indirect_Losses": ("Reported indirect loan losses; provider definition should be confirmed.", "Not currently used in dashboard KPI formulas."),
    "Total_Amount_of_Delinquent_Loans_Leases": ("Total dollar amount of delinquent loans and leases.", "Delinquency rate numerator and delinquent-loan balance KPI."),
    "Number_of_Delinquent_Loans_Leases": ("Number of delinquent loans and leases.", "Delinquent loan count KPI."),
    "Total_Amount_of_All_Loans_Charged_Off_YTD": ("Year-to-date total amount of loans charged off.", "Net charge-off and recovery-rate calculations; YTD period caveat."),
    "Total_Amount_of_All_Loans_Recovered_YTD": ("Year-to-date amount of loans recovered.", "Net charge-off and recovery-rate calculations; YTD period caveat."),
    "Equity_Acquired_in_Merger": ("Equity acquired through mergers.", "Net-worth component sum, subject to provider reconciliation."),
    "Equity_Undivided_Earnings_Amount": ("Undivided earnings component of equity.", "Net-worth component sum."),
    "Equity_Other_Reserves_Amount": ("Other reserves component of equity.", "Net-worth component sum."),
    "Equity_Appropriation_for_Non_Conforming_Investments": ("Equity appropriation for non-conforming investments.", "Net-worth component; verify inclusion/sign with source accounting definition."),
    "Equity_Noncontrolling_Interest_in_Consolidated_Subsidiaries": ("Noncontrolling interest in consolidated subsidiaries.", "Net-worth component sum; verify inclusion/sign."),
    "Equity_Accumulated_Unrealized_Net_Gains_Losses_on_Cash_Flow_Hedges": ("Accumulated unrealized net gains/losses on cash-flow hedges.", "Net-worth component sum; preserve sign."),
    "Equity_Other_Comprehensive_Income": ("Other comprehensive income component.", "Net-worth component sum; verify sign and inclusion."),
    "Equity_Accumulated_Unrealized_Gains_Losses_on_Available_for_Sale_Debt_Securities": ("Accumulated unrealized gains/losses on available-for-sale debt securities.", "Net-worth component sum; preserve negative values."),
    "Equity_Accumulated_Unrealized_Losses_for_OTTI": ("Accumulated unrealized losses for other-than-temporary impairment.", "Net-worth component sum; preserve sign."),
}

CATEGORY_FIELDS = {
    "Loans_by_Category_Unsecured_Credit_Card_Loans_Number": "Number of unsecured credit-card loans.",
    "Loans_by_Category_Unsecured_Credit_Card_Loans_Amount": "Balance of unsecured credit-card loans.",
    "Loans_by_Category_Payday_Alternative_Loans_Interest_Rate": "Reported interest rate for payday alternative loans.",
    "Loans_by_Category_Payday_Alternative_Loans_Number": "Number of payday alternative loans.",
    "Loans_by_Category_Payday_Alternative_Loans_Amount": "Balance of payday alternative loans.",
    "Loans_by_Category_Non_Federally_Guaranteed_Student_Loans_Interest_Rate": "Reported interest rate for non-federally guaranteed student loans.",
    "Loans_by_Category_Non_Federally_Guaranteed_Student_Loans_Number": "Number of non-federally guaranteed student loans.",
    "Loans_by_Category_Non_Federally_Guaranteed_Student_Loans_Amount": "Balance of non-federally guaranteed student loans.",
    "Loans_by_Category_All_Other_Unsecured_Loans_Interest_Rate": "Reported interest rate for all other unsecured loans.",
    "Loans_by_Category_All_Other_Unsecured_Loans_Number": "Number of all other unsecured loans.",
    "Loans_by_Category_All_Other_Unsecured_Loans_Amount": "Balance of all other unsecured loans.",
    "Loans_by_Category_New_Vehicle_Loans_Interest_Rate": "Reported interest rate for new-vehicle loans.",
    "Loans_by_Category_New_Vehicle_Loans_Number": "Number of new-vehicle loans.",
    "Loans_by_Category_New_Vehicle_Loans_Amount": "Balance of new-vehicle loans.",
    "Loans_by_Category_Used_Vehicle_Loans_Interest_Rate": "Reported interest rate for used-vehicle loans.",
    "Loans_by_Category_Used_Vehicle_Loans_Number": "Number of used-vehicle loans.",
    "Loans_by_Category_Used_Vehicle_Loans_Amount": "Balance of used-vehicle loans.",
    "Loans_by_Category_Leases_Receivable_Interest_Rate": "Reported interest rate for leases receivable.",
    "Loans_by_Category_Leases_Receivable_Number": "Number of leases receivable.",
    "Loans_by_Category_Leases_Receivable_Amount": "Balance of leases receivable.",
    "Loans_by_Category_All_Other_Secured_Non_Real_Estate_Loans_Interest_Rate": "Reported interest rate for other secured non-real-estate loans.",
    "Loans_by_Category_All_Other_Secured_Non_Real_Estate_Loans_Number": "Number of other secured non-real-estate loans.",
    "Loans_by_Category_All_Other_Secured_Non_Real_Estate_Loans_Amount": "Balance of other secured non-real-estate loans.",
    "Loans_by_Category_Loans_Secured_by_First_Lien_Interest_Rate": "Reported interest rate for first-lien loans.",
    "Loans_by_Category_Loans_Secured_by_First_Lien_Number": "Number of first-lien loans.",
    "Loans_by_Category_Loans_Secured_by_First_Lien_Amount": "Balance of first-lien loans.",
    "Loans_by_Category_Loans_Secured_by_Junior_Lien_Interest_Rate": "Reported interest rate for junior-lien loans.",
    "Loans_by_Category_Loans_Secured_by_Junior_Lien_Number": "Number of junior-lien loans.",
    "Loans_by_Category_Loans_Secured_by_Junior_Lien_Amount": "Balance of junior-lien loans.",
    "Loans_by_Category_Commercial_Loans_LoC_Real_Estate_Secured_Interest_Rate": "Reported interest rate for real-estate-secured commercial lines of credit.",
    "Loans_by_Category_Commercial_Loans_LoC_Real_Estate_Secured_Amount": "Duplicate source header; balance meaning is ambiguous until the provider identifies each physical column.",
    "Loans_by_Category_Commercial_Loans_LoC_Not_Real_Estate_Interest_Rate": "Reported interest rate for commercial lines not secured by real estate.",
    "Loans_by_Category_Commercial_Loans_LoC_Not_Real_Estate_Secured_Number": "Number of commercial lines not secured by real estate.",
    "Loans_by_Category_Commercial_Loans_LoC_Not_Real_Estate_Secured_Amount": "Balance of commercial lines not secured by real estate.",
    "Indirect_Loans_Number_of_Indirect_Loans_New_and_Used_Vehicle_Loans": "Number of indirect new- and used-vehicle loans.",
    "Indirect_Loans_Amount_of_Indirect_Loans_New_and_Used_Vehicle_Loans": "Balance of indirect new- and used-vehicle loans.",
    "Indirect_Loans_Number_of_Indirect_Loans_First_Lien_and_Junio_Lien_Residential_Loans": "Number of indirect first- and junior-lien residential loans.",
    "Indirect_Loans_Amount_of_Indirect_Loans_First_Lien_and_Junio_Lien_Residential_Loans": "Balance of indirect first- and junior-lien residential loans.",
    "Indirect_Loans_Number_of_Indirect_Loans_Commercial_Loans": "Number of indirect commercial loans.",
    "Indirect_Loans_Amount_of_Indirect_Loans_Commercial_Loans": "Balance of indirect commercial loans.",
    "Indirect_Loans_Number_of_Indirect_Loans_All_Other_Loans": "Number of other indirect loans.",
    "Indirect_Loans_Amount_of_Indirect_Loans_All_Other_Loans": "Balance of other indirect loans.",
}

CHARGEOFF_FIELDS = [
    "Charge_Off_Unsecured_Credit_Card_Loans_Amount", "Charge_Off_Leases_Receivable",
    "Charge_Off_Payday_Alternative_Loans", "Charge_Off_Non_Federally_Guaranteed_Student_Loans",
    "Charge_Off_New_Vehicle_Loans", "Charge_Off_Used_Vehicle_Loans", "Charge_Off_All_Other_Unsecured_Loans",
    "Charge_Off_All_Other_Secured_Non_Real_Estate_Loans", "Charge_Off_Loans_Secured_by_First_Lien",
    "Charge_Off_Loans_Secured_by_Junior_Lien", "Charge_Off_All_Other_Non_Commercial_Real_Estate_Loans",
    "Charge_Off_Construction_and_Development_Loans", "Charge_Off_Secured_by_Farmland",
    "Charge_Off_Secured_by_Multifamily", "Charge_Off_Secured_by_Owner_Occupied_Non_Farm_Non_Residential_Property",
    "Charge_Off_Secured_by_Non_Owner_Occupied_Non_Farm_Non_Residential_Property",
    "Charge_Off_Loans_to_finance_agricultural_production_other_loans_farmers",
    "Charge_Off_Commercial_and_Industrial_Loans", "Charge_Off_Unsecured_Commercial_Loans",
    "Charge_Off_Unsecured_Revolving_Lines_of_Credit_for_Commercial_Purposes",
]


def normalize_header(header):
    return re.sub(r"[^a-z0-9]", "", (header or "").lower())


def unique_headers(raw_headers):
    counts = Counter()
    values = []
    for index, header in enumerate(raw_headers, start=1):
        base = header.strip() or f"(blank header, physical column {index})"
        counts[base] += 1
        values.append((base, counts[base], index))
    return values


def load_source(path):
    with path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        headers = next(reader)
        first_row = next(reader, [])
    return headers, first_row


def data_type(name, sample):
    key = normalize_header(name)
    text = str(sample or "").strip()
    if not name:
        return "Unknown (blank header)"
    if key in {"quarter", "state", "cuname", "ceo", "website"} or key.endswith("name"):
        return "Text / category"
    if "date" in key:
        return "Date (source text)"
    if "rate" in key or "interestrate" in key:
        return "Percentage / rate (source numeric text)"
    if "number" in key or key in {"branches", "totalmembers", "cunumber", "joinnumber"}:
        return "Count / identifier (source text)"
    if "amount" in key or "assets" in key or "income" in key or "expenses" in key or "losses" in key or "shares" in key or text.startswith("$"):
        return "Currency amount (source text, dollars)"
    return "Text / numeric text"


def humanize(name):
    if not name:
        return "Unnamed source column; ask the provider for its meaning."
    words = name.replace("_", " ").strip()
    return words[:1].upper() + words[1:] + "."


def source_description(name, occurrence):
    if not name:
        return "No header supplied in source file; meaning cannot be determined from the CSV."
    if occurrence > 1:
        return "Duplicate header occurrence. The source gives this physical column the same name as another column; confirm its meaning with the provider before use."
    if name in SOURCE_DESCRIPTIONS:
        return SOURCE_DESCRIPTIONS[name][0]
    if name in CATEGORY_FIELDS:
        return CATEGORY_FIELDS[name]
    if name in CHARGEOFF_FIELDS:
        return f"Charge-off amount attributed to the loan category named in this source column; confirm whether this is YTD."
    if name.startswith("Loans_by_Category_"):
        return humanize(name.replace("Loans_by_Category_", "")) + " Loan-category field; amount, count or rate is indicated by the field name."
    return humanize(name)


def source_kpis(name, occurrence):
    if occurrence > 1:
        return "Not safely mapped: duplicate source column name; candidate for loan-category reporting only after provider clarification."
    if name in {"Quarter", "CYCLE_DATE"}:
        return "Period filters; quarter-over-quarter growth; trend comparison; all period-scoped KPIs."
    if name in {"CU_NUMBER", "CU_Name", "STATE"}:
        return "Institution identity; CU selection; rankings; peer filters; period matching."
    if name == "Total_Assets":
        return "Total assets; net-worth ratio; ROA; loan-to-assets; asset/member; asset market share/rank/growth; state map; asset buckets."
    if name == "Total_Members":
        return "Members; membership share/rank/percentile; members per branch; membership growth; state map."
    if name == "Branches":
        return "Branch count; members per branch; state map aggregation."
    if name == "Amount_of_Total_Loans_Leases":
        return "Total loans; loan-to-share; loan-to-assets; delinquency; charge-off ratio; balance-sheet ratios; loan growth."
    if name == "Total_Shares_Amount":
        return "Deposit/share balance; loan-to-share denominator; deposits/assets; deposit/share growth; market-share comparison."
    if name == "Number_Total_Shares":
        return "Share account count only; never use as a dollar denominator."
    if name == "NET_WORTH_AMOUNT" or name.startswith("Equity_"):
        return "Net worth / capital ratio; capital risk classification; competitive capital benchmark; gap analysis."
    if name == "Gross_Income":
        return "Efficiency ratio denominator; net income margin denominator; current portfolio-yield proxy."
    if name == "Operating_Expenses":
        return "Efficiency ratio numerator; competitive efficiency comparison."
    if name == "Net_Income":
        return "Net income KPI; ROA; ROE; net income margin; earnings alert; board summary."
    if name == "Total_Amount_of_Delinquent_Loans_Leases":
        return "Delinquency rate numerator; delinquent balance KPI; delinquency alert."
    if name == "Number_of_Delinquent_Loans_Leases":
        return "Delinquent loan count KPI."
    if name == "Total_Amount_of_All_Loans_Charged_Off_YTD":
        return "Net charge-off rate; charge-off amount; risk alert; recovery calculation."
    if name == "Total_Amount_of_All_Loans_Recovered_YTD":
        return "Net charge-off rate; recovery rate; recoveries amount; risk analysis."
    if name.startswith("Loans_by_Category_") or name.startswith("Indirect_Loans_"):
        return "Not used in current normalized dashboard; candidate source for loan product/indirect-loan mix after exact field validation."
    if name.startswith("Charge_Off_"):
        return "Not currently used; candidate category-level charge-off breakdown after field validation."
    if name in {"CEO", "Website", "JOIN_NUMBER", "Indirect_Losses"}:
        return "Institution profile / source reference only; no dashboard KPI formula currently uses this field."
    return "Not currently used in a dashboard KPI; see Notes for source quality or provider-definition questions."


def type_sample(name, rows, occurrence_index):
    for row in rows:
        if occurrence_index < len(row) and str(row[occurrence_index]).strip():
            return row[occurrence_index]
    return ""


def load_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        headers = next(reader)
        rows = list(reader)
    return headers, rows


def source_dictionary_rows():
    physical = defaultdict(dict)
    ordered_keys = []
    duplicate_names = set()
    for period, path in SOURCES.items():
        headers, rows = load_rows(path)
        duplicate_names.update(name for name, count in Counter(headers).items() if count > 1)
        occurrences = Counter()
        for index, raw_name in enumerate(headers):
            exact_name = raw_name.strip()
            key_name = exact_name if exact_name else ""
            occurrences[key_name] += 1
            key = (key_name, occurrences[key_name])
            if key not in physical:
                ordered_keys.append(key)
                physical[key] = {}
            physical[key][period] = {"index": index, "headers": headers, "rows": rows}
    output = []
    for name, occurrence in ordered_keys:
        source_data = physical[(name, occurrence)]
        q1 = source_data.get("Q1 2026")
        q2 = source_data.get("Q2 2026")
        present = []
        samples = []
        notes = []
        for period, record in (("Q1 2026", q1), ("Q2 2026", q2)):
            if record is None:
                continue
            position = record["index"] + 1
            present.append(f"{period}: column {position}")
            sample = type_sample(name, record["rows"], record["index"])
            samples.append(sample)
        if name == "":
            notes.append("Blank header in source; physical position and period listed. Do not map until provider supplies a name.")
        if name in duplicate_names:
            notes.append("Duplicate exact header; source positions are listed. Preserve as separate physical fields; resolve definitions with provider before calculations.")
        if name in {"NET_WORTH_AMOUNT"}:
            notes.append("Observed equal to Total_Assets for Affinity in Q1 and Q2 and for 3,986 Q1 records; validate before using as equity numerator.")
        if name in {"Total_Amount_of_All_Loans_Charged_Off_YTD", "Total_Amount_of_All_Loans_Recovered_YTD"}:
            notes.append("YTD basis; quarter-over-quarter numerator requires Q2 YTD minus Q1 YTD and consistent definitions.")
        if name == "Number_Total_Shares":
            notes.append("Share account count, not a dollar amount; never use as loan-to-share denominator.")
        if name == "Total_Shares_Amount":
            notes.append("Dollar share balance; correct loan-to-share denominator.")
        if name == "CYCLE_DATE":
            notes.append("Source date is displayed as DD-MM-YYYY in files reviewed.")
        description = source_description(name, 2 if name in duplicate_names else occurrence)
        dtype = data_type(name, samples[0] if samples else "")
        used_in_kpis = "Not reliably mapped: duplicate source column name; resolve with provider before using this field."
        if name not in duplicate_names:
            used_in_kpis = source_kpis(name, occurrence)
        output.append((
            name,
            description,
            used_in_kpis,
            dtype,
            "; ".join(present + notes),
        ))
    return output


def kpi_dictionary_rows():
    return KPI_ROWS


def style_sheet(sheet, widths, table_name):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="C9A227"))
    sheet.row_dimensions[1].height = 36
    for row_index in range(2, sheet.max_row + 1):
        for cell in sheet[row_index]:
            cell.font = Font(name="Aptos", size=9, color="243242")
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if row_index % 2 == 0 else "F3F6F9")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="D8E0E8"))
        sheet.row_dimensions[row_index].height = 76 if sheet.title == "KPI Data Dictionary" else 54
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)


def main():
    missing = [str(path) for path in SOURCES.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Source CSV not found: " + "; ".join(missing))
    workbook = Workbook()
    kpi_sheet = workbook.active
    kpi_sheet.title = "KPI Data Dictionary"
    kpi_sheet.append([
        "KPI Name", "Derived or Direct Source Column?", "Source Column Name(s)",
        "Business Logic / Formula", "Plain English Explanation", "Expected Range (If Applicable)",
        "Notes / Assumptions", "Current Dashboard Logic", "Recommended Logic",
    ])
    for row in kpi_dictionary_rows():
        kpi_sheet.append(row)
    style_sheet(kpi_sheet, {"A": 34, "B": 25, "C": 65, "D": 68, "E": 50, "F": 38, "G": 58, "H": 62, "I": 66}, "KpiDictionary")

    source_sheet = workbook.create_sheet("Source Column Dictionary")
    source_sheet.append([
        "Source Column Name", "Source Data Description", "Used In KPI(s)", "Data Type", "Notes",
    ])
    source_rows = source_dictionary_rows()
    for row in source_rows:
        source_sheet.append(row)
    style_sheet(source_sheet, {"A": 76, "B": 58, "C": 64, "D": 36, "E": 74}, "SourceColumnDictionary")
    workbook.properties.title = "CU360 Source-to-KPI Data Dictionary"
    workbook.properties.subject = "Business mapping from Q1/Q2 2026 NCUA CSV fields to dashboard KPIs"
    workbook.properties.creator = "CU360 Intelligence"
    workbook.save(OUTPUT)

    check = load_workbook(OUTPUT, read_only=True, data_only=True)
    print(f"Created {OUTPUT}")
    print(f"Sheets: {check.sheetnames}")
    print(f"KPI entries: {check['KPI Data Dictionary'].max_row - 1}")
    print(f"Physical source-column entries: {check['Source Column Dictionary'].max_row - 1}")
    print(f"Q1/Q2 periods included in source dictionary: {', '.join(SOURCES)}")
    check.close()


if __name__ == "__main__":
    main()
